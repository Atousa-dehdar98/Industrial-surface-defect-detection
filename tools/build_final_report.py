from pathlib import Path
import json

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Industrial_Surface_Defect_Detection_Final_Report.docx"
FROZEN = json.loads((ROOT / "outputs/analysis/milestone3/validation_threshold_selection_summary.json").read_text())
FINE = json.loads((ROOT / "outputs/analysis/milestone5/fine_tuned_threshold_selection_summary.json").read_text())
TEST = json.loads((ROOT / "outputs/analysis/milestone7/final_test_evaluation_summary.json").read_text())
GRADCAM = json.loads((ROOT / "outputs/analysis/milestone6/gradcam_summary.json").read_text())

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT = "F2F4F7"
MUTED = "666666"
RED = "9B1C1C"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_fixed_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for cell, value in zip(hdr.cells, headers):
        cell.text = str(value)
        set_cell_shading(cell, LIGHT)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(9)
    for row_values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values):
            cell.text = str(value)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(9)
    set_fixed_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def add_picture(doc, path, width, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(1)
section.left_margin = section.right_margin = Inches(1)
section.header_distance = section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10
for style_name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK_BLUE, 8, 4),
):
    style = styles[style_name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
for name in ("List Bullet", "List Number"):
    styles[name].font.name = "Calibri"
    styles[name].font.size = Pt(11)
    styles[name].paragraph_format.space_after = Pt(8)
    styles[name].paragraph_format.line_spacing = 1.167

header = section.header.paragraphs[0]
header.text = "FINAL PROJECT REPORT  |  DSPT 0326"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in header.runs:
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("Industrial Surface Defect Detection | August 2026")
footer_run.font.size = Pt(8.5)
footer_run.font.color.rgb = RGBColor.from_string(MUTED)

# Cover: editorial report treatment.
doc.add_paragraph().paragraph_format.space_after = Pt(72)
kicker = doc.add_paragraph()
kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = kicker.add_run("DATA SCIENCE BOOTCAMP - FINAL PROJECT")
r.bold = True
r.font.size = Pt(11)
r.font.color.rgb = RGBColor.from_string(BLUE)
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(12)
title.paragraph_format.space_after = Pt(10)
r = title.add_run("Industrial Surface Defect Detection\nUsing Deep Learning and Explainable AI")
r.bold = True
r.font.size = Pt(25)
r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("Supervised Binary Classification of MVTec AD Tile Images")
r.font.size = Pt(14)
r.font.color.rgb = RGBColor.from_string(MUTED)
doc.add_paragraph().paragraph_format.space_after = Pt(80)
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("Prepared by: Atousa\n").bold = True
meta.add_run("Cohort: DSPT 0326\nFinal submission: August 26, 2026\nPresentation: August 27, 2026")
doc.add_page_break()

doc.add_heading("Executive Summary", level=1)
doc.add_paragraph(
    "This project develops a reproducible supervised classifier for Good versus Defective tile surfaces. "
    "A custom CNN was compared with ImageNet-pretrained MobileNetV2, followed by controlled fine-tuning, validation-only threshold calibration, Grad-CAM analysis, and a confidence-based human-review mechanism."
)
lead = doc.add_paragraph()
lead.paragraph_format.space_before = Pt(6)
lead.paragraph_format.space_after = Pt(10)
run = lead.add_run(
    "Key result: the locked final test achieved 94.12% accuracy, 90.91% defective precision, 83.33% defective recall, 86.96% F1, 97.44% ROC-AUC, and 94.44% PR-AUC."
)
run.bold = True
run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
doc.add_paragraph(
    "Two defective test images were classified as Good. This safety-relevant limitation is retained as observed; neither the model nor thresholds were modified after test evaluation."
)

doc.add_heading("1. Problem Definition and Objectives", level=1)
doc.add_paragraph(
    "Surface defects can create quality, cost, and safety risks in industrial production. Manual visual inspection is resource-intensive and can be inconsistent. The primary objective was to classify tile images as Good or Defective while preserving transparency about uncertainty and failure cases."
)
add_bullet(doc, "Build a leakage-safe supervised binary classification workflow.")
add_bullet(doc, "Compare a custom CNN with transfer learning and controlled fine-tuning.")
add_bullet(doc, "Prioritize defective recall alongside precision, F1, ROC-AUC, and PR-AUC.")
add_bullet(doc, "Use Grad-CAM for spatial sensitivity inspection.")
add_bullet(doc, "Implement Good, Defective, and Needs Human Review decisions.")

doc.add_heading("2. Dataset and Scientific Framing", level=1)
doc.add_paragraph(
    "The project uses the Tile category of MVTec AD. MVTec AD was designed for unsupervised anomaly detection, and its original training folder generally contains only Good images. This project intentionally reformulates the data as supervised binary classification. That choice enables the course objective but differs from the benchmark's original protocol."
)
add_table(doc, ["Split", "Images", "Good", "Defective", "Purpose"], [
    ["Train", "242", "184", "58", "Parameter learning"],
    ["Validation", "54", "40", "14", "Model and threshold selection"],
    ["Test", "51", "39", "12", "One-time final reporting"],
], [1200, 1000, 1000, 1100, 5060])
doc.add_paragraph(
    "Assignments were made at the original-image level with seed 42. Group identifiers and SHA-256 hashes were checked to prevent cross-split leakage. Augmentation was applied only to training batches."
)

doc.add_heading("3. Data Preparation and Exploratory Analysis", level=1)
doc.add_paragraph(
    "The data audit verified file availability, image decoding, class labels, split distributions, and duplicate isolation. Images were resized to 224 x 224 RGB tensors. Training augmentation included flips, small rotations, zoom, translation, and contrast variation. Validation and test images received resizing only. Class weights addressed imbalance during training."
)

doc.add_heading("4. Modeling Strategy", level=1)
doc.add_heading("4.1 Custom CNN baseline", level=2)
doc.add_paragraph(
    "The baseline used three convolutional blocks with BatchNormalization, ReLU, and max pooling, followed by global average pooling, dropout, and a sigmoid classifier. It provided an interpretable reference but overfit the small dataset and produced weak validation discrimination."
)
doc.add_heading("4.2 Transfer learning", level=2)
doc.add_paragraph(
    "MobileNetV2 with ImageNet weights was used as a frozen feature extractor. A global average pooling layer, dropout, and a one-unit sigmoid head were trained with binary cross-entropy. This addressed the instructor's recommendation to use transfer learning for the relatively small dataset."
)
doc.add_heading("4.3 Controlled fine-tuning", level=2)
doc.add_paragraph(
    "After a successful smoke test, the last 20 MobileNetV2 layers were made eligible for training while BatchNormalization layers remained frozen. Adam used a learning rate of 1e-5. EarlyStopping, ModelCheckpoint, ReduceLROnPlateau, and TerminateOnNaN controlled the run. Training stopped after 9 epochs; epoch 5 had the best validation loss."
)
add_picture(doc, ROOT / "outputs/figures/milestone4/mobilenetv2_finetuned_standard_learning_curves.png", 6.3, "Figure 1. Executed fine-tuning learning curves. The best validation loss occurred at epoch 5.")

doc.add_heading("5. Validation Results and Model Selection", level=1)
add_table(doc, ["Model", "Accuracy", "Precision", "Recall", "ROC-AUC", "PR-AUC"], [
    ["Custom CNN", "25.93%", "25.93%", "100.00%", "0.761", "0.446"],
    ["Frozen MobileNetV2", "88.89%", "83.33%", "71.43%", "0.975", "0.937"],
    ["Fine-tuned MobileNetV2", "94.44%", "92.31%", "85.71%", "0.989", "0.972"],
], [2800, 1250, 1350, 1200, 1300, 1460])
doc.add_paragraph(
    "Fine-tuned MobileNetV2 was selected because it improved validation loss and ranking metrics while remaining computationally efficient. At the validation-selected binary threshold of 0.152666, recall reached 100%, precision 82.35%, F1 90.32%, and accuracy 94.44%."
)

doc.add_heading("6. Confidence-Based Human Review", level=1)
doc.add_paragraph(
    "The binary threshold maximized validation F1 among candidates with defective recall of at least 90%. The lower boundary was fixed at 0.152666 and the automatic-Defective boundary at 0.439. Predictions between the boundaries were sent to review."
)
add_table(doc, ["Decision", "Probability rule", "Validation count", "Final test count"], [
    ["Good", "p < 0.152666", "37", "40"],
    ["Needs Human Review", "0.152666 <= p < 0.439", "4", "1"],
    ["Defective", "p >= 0.439", "13", "10"],
], [2200, 3300, 1900, 1960])
doc.add_paragraph(
    "Validation suggested zero defects would be automatically accepted as Good. On the locked test set, two defects fell below the Good boundary. The policy therefore reduced review volume but did not provide a general safety guarantee."
)

doc.add_heading("7. Grad-CAM Explainability", level=1)
doc.add_paragraph(
    f"Grad-CAM was computed at layer {GRADCAM['target_layer']} for {GRADCAM['selected_examples']} validation examples: correct Good and Defective predictions, all False Positives, and low-confidence or review cases. Crack images showed attention near visible crack structures. False positives often showed broad attention over texture or borders."
)
add_picture(doc, ROOT / "outputs/figures/milestone6/gradcam/fine_tuned_validation_gradcam_panel.png", 2.30, "Figure 2. Grad-CAM validation panel. Red and yellow areas indicate stronger gradient-based sensitivity for the Defective output.")
doc.add_paragraph(
    "Grad-CAM is a coarse sensitivity visualization. It does not establish a causal explanation and may respond to texture, lighting, borders, or correlated artifacts."
)

doc.add_heading("8. Final One-Time Test Evaluation", level=1)
metrics = TEST["binary_metrics"]
add_table(doc, ["Metric", "Final test value"], [
    ["Accuracy", f"{metrics['accuracy']:.2%}"],
    ["Defective precision", f"{metrics['precision']:.2%}"],
    ["Defective recall", f"{metrics['recall']:.2%}"],
    ["F1-score", f"{metrics['f1']:.2%}"],
    ["ROC-AUC", f"{metrics['roc_auc']:.3f}"],
    ["PR-AUC", f"{metrics['pr_auc']:.3f}"],
    ["TN / FP / FN / TP", f"{metrics['tn']} / {metrics['fp']} / {metrics['fn']} / {metrics['tp']}"],
], [5000, 4360])
add_picture(doc, ROOT / "outputs/figures/milestone7/final_test_metrics.png", 6.3, "Figure 3. Locked final test confusion matrix, ROC curve, and precision-recall curve.")
doc.add_paragraph(
    "Per-defect recall was 100% for crack, glue_strip, and oil, and 50% for gray_stroke and rough. Because each subgroup contains only two or three test images, these estimates have high uncertainty."
)
add_picture(doc, ROOT / "outputs/figures/milestone7/final_test_error_examples.png", 6.3, "Figure 4. All final test errors: two False Negatives and one False Positive. The False Positive was referred for human review.")

doc.add_heading("9. Limitations and Practical Risks", level=1)
add_bullet(doc, "The defective validation and test sample sizes are small, producing uncertain estimates.")
add_bullet(doc, "The supervised reformulation differs from the original MVTec AD anomaly-detection protocol.")
add_bullet(doc, "Validation-based thresholds did not prevent two test defects from being automatically marked Good.")
add_bullet(doc, "Subgroup performance was weakest for gray_stroke and rough defects.")
add_bullet(doc, "Grad-CAM is not a causal explanation and cannot replace domain inspection.")
add_bullet(doc, "The prototype has not been tested for production latency, drift, lighting changes, camera changes, or calibration stability.")

doc.add_heading("10. Ethical and Operational Considerations", level=1)
doc.add_paragraph(
    "Automated inspection should support rather than silently replace accountable quality-control processes. False negatives can allow defective products to pass, while false positives increase rework and inspection cost. Deployment would require a documented escalation policy, monitoring by defect type, audit logs, operator training, and periodic reassessment with new data."
)

doc.add_heading("11. Conclusion", level=1)
doc.add_paragraph(
    "Transfer learning was appropriate for the limited dataset and materially outperformed the custom CNN. Controlled fine-tuning produced the strongest validation model and strong test ranking performance. Nevertheless, final-test recall of 83.33% and two automatically accepted defects show that the current system is an experimental prototype rather than a production-ready inspection solution. The project demonstrates a reproducible data-science lifecycle, transparent error reporting, explainability diagnostics, and disciplined separation between validation decisions and final test reporting."
)

doc.add_heading("12. Recommended Next Development Cycle", level=1)
add_bullet(doc, "Collect more representative gray_stroke, rough, and borderline Good images.")
add_bullet(doc, "Create a new untouched holdout set before any further model or threshold changes.")
add_bullet(doc, "Evaluate calibration methods and review thresholds under an explicit cost model.")
add_bullet(doc, "Test robustness to lighting, camera, crop, and production-line variation.")
add_bullet(doc, "Run a prospective human-in-the-loop pilot with quality-control experts.")

doc.add_heading("References", level=1)
references = doc.add_paragraph()
references.paragraph_format.space_after = Pt(0)
references.add_run(
    "Bergmann, P., Fauser, M., Sattlegger, D., & Steger, C. (2019). MVTec AD - A Comprehensive Real-World Dataset for Unsupervised Anomaly Detection. Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition.\n"
    "Data Science Bootcamp, DSPT 0326. Final Project Description and Milestone documents, 2026.\n"
    "Keras Applications documentation: MobileNetV2 architecture and ImageNet preprocessing."
)

doc.save(OUT)
print(OUT)
