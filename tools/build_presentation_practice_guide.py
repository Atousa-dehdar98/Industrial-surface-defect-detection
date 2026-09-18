from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = r"D:\Final Project\Industrial_Surface_Defect_Detection_Presentation_Practice_Guide.docx"

BLUE = "2E74B5"
DARK = "1F4D78"
NAVY = "0B2545"
MUTED = "667085"
LIGHT = "E8EEF5"
CALLOUT = "F4F6F9"
GOLD = "7A5A00"
RED = "9B1C1C"

slides = [
    {
        "n": 1, "title": "Industrial Surface Defect Detection", "time": "40–50 seconds",
        "message": "پروژه را در یک جمله معرفی کنید: تشخیص دودویی عیب روی Tile همراه با توضیح‌پذیری و ارجاع موارد نامطمئن به انسان.",
        "script": "Good morning, and thank you for the opportunity to present my final project. This project focuses on industrial surface defect detection using the Tile category of the MVTec Anomaly Detection dataset. I framed the task as supervised binary image classification: each tile is classified as either good or defective. Because the dataset is relatively small, I used transfer learning with MobileNetV2, followed by controlled fine-tuning. I also used Grad-CAM to inspect what the model attends to, and I added a confidence-based human-review policy so that uncertain predictions are not treated as fully automatic decisions. My goal was not only strong accuracy, but also a workflow that is transparent about risk and model limitations.",
        "transition": "I will begin with the practical problem and the main risks that shaped the methodology.",
        "cue": "MVTec Tile • binary classification • transfer learning • Grad-CAM • human review"
    },
    {
        "n": 2, "title": "Problem, Objective, and Risk", "time": "50–60 seconds",
        "message": "نشان دهید که مسئله فقط Accuracy نیست؛ خطای از دست دادن عیب مهم‌تر است و محدودیت داده کوچک نیز باید مدیریت شود.",
        "script": "The practical objective is to separate acceptable tiles from defective tiles. In an industrial setting, the most important error is usually a false negative: a defective tile that the system labels as good. For that reason, I monitored defective-class recall and the confusion matrix, not accuracy alone. The dataset also creates two methodological challenges. First, it is small and imbalanced, so a model trained from scratch can overfit or learn unstable decision rules. Second, image leakage between train, validation, and test sets could make the results look unrealistically strong. These risks motivated three design choices: leakage-aware splitting, transfer learning, and validation-based threshold selection before the test set was opened.",
        "transition": "The next slide shows how the data was separated to protect the evaluation.",
        "cue": "False negatives • small data • imbalance • leakage • validation first"
    },
    {
        "n": 3, "title": "Dataset and Leakage-Aware Split", "time": "55–65 seconds",
        "message": "تعداد نمونه‌ها و استقلال splitها را دقیق بگویید؛ augmentation فقط روی train انجام شد.",
        "script": "After validating the dataset structure, I created three fixed and non-overlapping partitions. The training set contains 242 images: 184 good and 58 defective. The validation set contains 54 images: 40 good and 14 defective. The final test set contains 51 images: 39 good and 12 defective. I checked for duplicate file paths and content overlap across the partitions to reduce leakage risk. All threshold selection, model comparison, and early stopping decisions used the validation set. The test set remained untouched until the final evaluation. Data augmentation was applied only to training images, while validation and test images used deterministic preprocessing. This separation is essential because otherwise augmentation or repeated images could contaminate the reported performance.",
        "transition": "With the evaluation protocol fixed, I compared learning from scratch with transfer learning.",
        "cue": "242 / 54 / 51 • no overlap • augmentation train only • test locked"
    },
    {
        "n": 4, "title": "Model Development Journey", "time": "55–65 seconds",
        "message": "مسیر سه‌مرحله‌ای را توضیح دهید و تاکید کنید که transfer learning پاسخ مستقیم به کوچک بودن دیتاست بود.",
        "script": "I evaluated three stages of model development. The first was a small custom convolutional neural network trained from scratch. It served as a baseline, but the validation behavior showed that it was not learning a useful balanced decision boundary. The second stage used MobileNetV2 with ImageNet-pretrained features and a frozen backbone. This immediately improved generalization, which supports the instructor's recommendation to use transfer learning for a small dataset. The third stage was controlled fine-tuning. I kept most of the pretrained network frozen and allowed only the final part of the backbone to adapt to tile textures. This staged approach made it possible to measure the contribution of pretrained features before introducing additional flexibility.",
        "transition": "The validation comparison makes the benefit of this staged approach clear.",
        "cue": "Custom CNN → frozen MobileNetV2 → controlled fine-tuning"
    },
    {
        "n": 5, "title": "Validation Model Comparison", "time": "70–80 seconds",
        "message": "اعداد را آرام بگویید و خطای تفسیر Recall صددرصدی baseline را توضیح دهید.",
        "script": "At the default probability threshold of zero point five, the custom CNN achieved only 25.93 percent accuracy and precision, although its defective recall was 100 percent. That recall is misleading in isolation because the model predicted too many images as defective. Its ROC-AUC was 0.761 and PR-AUC was 0.446. The frozen MobileNetV2 improved accuracy to 88.89 percent, with 83.33 percent precision, 71.43 percent recall, and a PR-AUC of 0.937. After controlled fine-tuning, accuracy reached 94.44 percent, precision 92.31 percent, recall 85.71 percent, ROC-AUC 0.989, and PR-AUC 0.972. Therefore, fine-tuning produced the strongest and most balanced validation model.",
        "transition": "I fine-tuned conservatively to preserve the useful pretrained representation.",
        "cue": "CNN 25.93% • frozen 88.89% • fine-tuned 94.44% • PR-AUC .972"
    },
    {
        "n": 6, "title": "Controlled Fine-Tuning", "time": "50–60 seconds",
        "message": "کنترل ریسک overfitting را توضیح دهید: فقط ۲۰ لایه پایانی، BatchNorm ثابت، learning rate کم و early stopping.",
        "script": "For fine-tuning, I made only the last 20 MobileNetV2 layers eligible for training. Batch-normalization layers remained frozen because updating their statistics with small batches can destabilize pretrained features. I used the Adam optimizer with a learning rate of one times ten to the minus five. Training stopped after nine epochs, and the best validation checkpoint came from epoch five. The complete fine-tuning run took approximately one minute and 56 seconds on the available environment. This was intentionally a controlled experiment rather than an extensive hyperparameter search. It added limited domain adaptation while reducing the risk of catastrophic forgetting and overfitting on a small dataset.",
        "transition": "After selecting the model, I calibrated the decision policy on validation data.",
        "cue": "Last 20 layers • BatchNorm frozen • Adam 1e-5 • best epoch 5 • 1m56s"
    },
    {
        "n": 7, "title": "Validation-Locked Decision Policy", "time": "70–80 seconds",
        "message": "تفاوت threshold طبقه‌بندی و بازه human review را روشن کنید و تاکید کنید test در انتخاب نقشی نداشت.",
        "script": "A probability of zero point five is not automatically the best operating threshold when false negatives have a higher cost. Using only validation predictions, I selected a binary threshold of 0.152666. At that operating point, validation defective recall was 100 percent, precision was 82.35 percent, and F1 was 90.32 percent. I then defined a confidence-based review region. Probabilities below 0.152666 were automatically labeled good, probabilities above 0.439 were automatically labeled defective, and the values in between were sent for human review. This policy produced a validation review rate of 7.41 percent. Most importantly, both thresholds were locked before final test evaluation; I did not adjust them after seeing test errors.",
        "transition": "Before the final test, Grad-CAM provided a qualitative check of the model's attention.",
        "cue": "Low .152666 • high .439 • val recall 100% • review 7.41% • locked"
    },
    {
        "n": 8, "title": "Grad-CAM Explainability", "time": "55–65 seconds",
        "message": "Grad-CAM را ابزار تشخیصی معرفی کنید، نه اثبات علت تصمیم مدل.",
        "script": "Grad-CAM was used as a qualitative diagnostic tool. It produces a heatmap that highlights spatial regions associated with the model's output. I examined correctly classified examples, false positives, and low-confidence cases in the explainability notebook; the presentation emphasizes challenging review examples. This helps check whether attention is concentrated near visible surface irregularities or on irrelevant background patterns. However, I do not interpret Grad-CAM as a causal explanation or proof that the model has learned the correct physical concept. It is better treated as supporting evidence for inspection, error analysis, and human review, especially when combined with the predicted probability and the original image.",
        "transition": "After locking the model and policy, I evaluated them once on the untouched test set.",
        "cue": "Qualitative diagnostic • attention regions • not causal • support review"
    },
    {
        "n": 9, "title": "Final Test Performance", "time": "70–80 seconds",
        "message": "نتیجه واقعی test را گزارش کنید و confusion matrix را با صدای واضح بخوانید.",
        "script": "On the untouched final test set, the fine-tuned model achieved 94.12 percent accuracy. Defective-class precision was 90.91 percent, recall was 83.33 percent, and F1 was 86.96 percent. The ranking metrics remained strong, with ROC-AUC of 0.974 and PR-AUC of 0.944. The confusion matrix contained 38 true negatives, one false positive, two false negatives, and 10 true positives. These are strong results for the available sample, but the two missed defects matter operationally. The gap between 100 percent validation recall at the selected threshold and 83.33 percent test recall also shows why validation performance should not be presented as a guarantee of future performance.",
        "transition": "The error analysis shows exactly where the remaining operational weakness lies.",
        "cue": "Acc 94.12 • Prec 90.91 • Recall 83.33 • AUC .974 • TN38 FP1 FN2 TP10"
    },
    {
        "n": 10, "title": "Error Analysis and Human Review", "time": "75–90 seconds",
        "message": "دو FN و یک FP را دقیق بگویید؛ سیاست review فقط FP را گرفت و FNها را نگرفت.",
        "script": "The two false negatives came from the gray-stroke and rough defect categories. Their defective probabilities were 0.119 and 0.047, so both fell below the lower threshold and were automatically labeled good. Test recall was 100 percent for crack, glue-strip, and oil defects, but only 50 percent for gray-stroke and rough defects. The single false positive had a probability of 0.192, which placed it inside the human-review interval. Therefore, the review policy routed one of 51 test images, or 1.96 percent, to a person and provided 98.04 percent automatic coverage. But it did not catch the two false negatives. This is an important limitation: confidence-based review reduces some uncertainty, but confidence alone does not guarantee safety.",
        "transition": "I will conclude with what the project demonstrates and what should be improved next.",
        "cue": "FN gray_stroke .119 • FN rough .047 • FP .192 reviewed • review 1.96% • limitation"
    },
    {
        "n": 11, "title": "Conclusion and Next Steps", "time": "55–65 seconds",
        "message": "پروژه را موفق اما غیر-production-ready جمع‌بندی کنید و قدم بعدی را مشخص بگویید.",
        "script": "In conclusion, transfer learning was clearly more effective than training a small CNN from scratch, and controlled fine-tuning produced the best validation and test results. The final model achieved strong discrimination and 94.12 percent test accuracy, while Grad-CAM added a useful qualitative inspection layer. The project also demonstrates the value of separating model prediction from operational decision policy. However, the remaining false negatives show that this is not yet a production-ready safety system. The next steps would be to collect more examples of subtle gray-stroke and rough defects, calibrate uncertainty on new validation data, consider class- or category-aware review rules, and confirm all improvements on a new untouched holdout set. Thank you, and I welcome your questions.",
        "transition": "Thank you. I am ready for your questions.",
        "cue": "Transfer learning wins • explainability • policy layer • not production-ready • new holdout"
    },
]

questions = [
    ("Why did you formulate an anomaly-detection dataset as supervised binary classification?", "The project scope required binary supervised classification. I used the available good and defect labels to learn a decision boundary, while preserving defect subtypes for stratified analysis rather than predicting them as separate classes."),
    ("How did you prevent data leakage?", "I created fixed train, validation, and test partitions, checked for overlap, applied augmentation only to training data, and kept the test set closed until the model and thresholds were locked."),
    ("Why did you choose MobileNetV2?", "It provides ImageNet-pretrained visual features with relatively low computational cost. That makes it suitable for a small dataset and for reproducible experiments on limited hardware."),
    ("Why not train a deeper model from scratch?", "The custom CNN baseline generalized poorly. With limited data, a larger model from scratch would increase overfitting risk without evidence that it would improve generalization."),
    ("Why did you fine-tune only the last 20 layers?", "I wanted limited adaptation to tile textures while preserving most pretrained features. The choice was a controlled experiment, not the result of a large tuning search."),
    ("Why were batch-normalization layers frozen?", "Their running statistics can become unstable with small datasets and batches. Freezing them preserves the pretrained normalization behavior during controlled fine-tuning."),
    ("Why is the classification threshold lower than 0.5?", "The threshold was selected on validation data to prioritize defective recall. A default threshold is arbitrary when error costs are asymmetric."),
    ("Why was validation recall 100 percent but test recall only 83.33 percent?", "The threshold fit the finite validation sample. The test set contains new variation, and two subtle defects received low confidence. This is normal generalization uncertainty and a reason not to tune on test data."),
    ("Did the human-review policy solve the false-negative problem?", "No. It routed the one false positive to review, but both false negatives were confidently below the lower threshold. This is a key limitation and is reported explicitly."),
    ("Why report PR-AUC as well as ROC-AUC?", "The defective class is the minority class. PR-AUC focuses on the precision-recall trade-off for that positive class and is often more informative under imbalance."),
    ("How can the baseline have 100 percent recall and still be poor?", "It predicted too many samples as defective. Recall alone ignores false positives, so its low precision and accuracy reveal the degenerate behavior."),
    ("What does Grad-CAM prove?", "It does not prove causality. It provides a qualitative view of spatial attention that can reveal plausible focus or suspicious shortcuts and can support human inspection."),
    ("Why did you not perform extensive hyperparameter tuning?", "The dataset is small, broad tuning could overfit the validation set, and long-running tuning was outside the controlled project scope. I prioritized a transparent staged comparison."),
    ("Is the system ready for production?", "No. The test sample is small and two defects were missed. Production use would require more representative data, calibration, monitoring, process-specific cost analysis, and a new untouched evaluation."),
    ("What is the most important next improvement?", "Collect more representative gray-stroke and rough examples, then redesign and validate the review policy on new validation data and confirm it on a fresh holdout set."),
]

short_script = [
    "Slide 1 — This project detects good versus defective MVTec Tile images using supervised binary classification. Because the dataset is small, I used transfer learning, Grad-CAM, and a confidence-based human-review policy.",
    "Slide 2 — The main operational risk is a false negative: a defective tile labeled good. Therefore, I evaluated recall and the confusion matrix, not accuracy alone, and I treated leakage and overfitting as major methodological risks.",
    "Slide 3 — The fixed split contains 242 training, 54 validation, and 51 test images. I checked overlap, augmented only training data, used validation for all choices, and kept test data untouched until the end.",
    "Slide 4 — I compared a custom CNN, frozen MobileNetV2 transfer learning, and controlled fine-tuning. The pretrained backbone generalized much better than learning from scratch.",
    "Slide 5 — Fine-tuning produced the best validation result: 94.44 percent accuracy, 85.71 percent recall, and PR-AUC of 0.972. The baseline's 100 percent recall was misleading because it over-predicted defects.",
    "Slide 6 — I fine-tuned only the last 20 layers, froze batch normalization, and used a learning rate of one times ten to the minus five. The best checkpoint was epoch five.",
    "Slide 7 — On validation data, I selected 0.152666 as the lower decision threshold and 0.439 as the upper review threshold. Values between them were sent for human review. Both thresholds were locked before test evaluation.",
    "Slide 8 — Grad-CAM was used as a qualitative diagnostic to inspect attention on correct, incorrect, and uncertain examples. It supports inspection, but it is not a causal explanation.",
    "Slide 9 — Final test accuracy was 94.12 percent, defective recall was 83.33 percent, ROC-AUC was 0.974, and PR-AUC was 0.944. The confusion matrix was 38 true negatives, one false positive, two false negatives, and 10 true positives.",
    "Slide 10 — The two missed defects were gray-stroke and rough, and both were automatically labeled good. The review policy routed the false positive, but not the false negatives, so confidence alone was not sufficient for safety.",
    "Slide 11 — Transfer learning and controlled fine-tuning were effective, but the system is not production-ready. The next priority is more subtle-defect data, improved uncertainty handling, and confirmation on a new untouched holdout set.",
]

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_width(cell, dxa):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(dxa)); tcW.set(qn("w:type"), "dxa")

def set_table_geometry(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    tblW.set(qn("w:w"), "9360"); tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd"); tblPr.append(tblInd)
    tblInd.set(qn("w:w"), "120"); tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(width)); grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells): set_cell_width(cell, widths[i])

def set_font(run, size=None, color=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Arial")
    if size: run.font.size = Pt(size)
    if color: run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic

def rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = pPr.find(qn("w:bidi"))
    if bidi is None: bidi = OxmlElement("w:bidi"); pPr.append(bidi)
    bidi.set(qn("w:val"), "1")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

def add_text(doc, text, *, bold=False, italic=False, color=None, size=11, after=6, rtl_text=False, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.keep_together = keep
    r = p.add_run(text); set_font(r, size=size, color=color, bold=bold, italic=italic)
    if rtl_text: rtl(p)
    return p

def add_callout(doc, label, text, fill=CALLOUT, rtl_text=False):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.15
    r = p.add_run(label + " "); set_font(r, size=10.5, bold=True, color=DARK)
    r = p.add_run(text); set_font(r, size=10.5)
    if rtl_text: rtl(p)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page "); set_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"; normal.font.size = Pt(11); normal.font.color.rgb = RGBColor.from_string("202124")
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.paragraph_format.space_before = Pt(0); normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.25
for lvl, size, color, before, after in [(1,16,BLUE,18,10),(2,13,BLUE,14,7),(3,12,DARK,10,5)]:
    st = styles[f"Heading {lvl}"]; st.font.name="Calibri"; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True

header = sec.header.paragraphs[0]
header.text = "PRESENTATION PRACTICE GUIDE  |  MVTec AD Tile"
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
for r in header.runs: set_font(r, size=8.5, color=MUTED, bold=True)
add_page_number(sec.footer.paragraphs[0])

# Editorial cover
add_text(doc, "FINAL PROJECT • ORAL PRESENTATION", bold=True, color=GOLD, size=10, after=24)
for _ in range(3): add_text(doc, "", after=14)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(10)
r = p.add_run("Industrial Surface Defect Detection"); set_font(r, size=29, color=NAVY, bold=True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(8)
r = p.add_run("Presentation Practice Guide"); set_font(r, size=18, color=BLUE, bold=True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(28)
r = p.add_run("MVTec AD Tile • Transfer Learning • Grad-CAM • Confidence-Based Human Review"); set_font(r, size=11.5, color=MUTED)
add_callout(doc, "Target duration:", "10–12 minutes  |  Full English script + Persian coaching + cue cards + Q&A + 5-minute backup", fill=LIGHT)
add_text(doc, "Prepared as a rehearsal companion to the final presentation deck and report.", italic=True, color=MUTED, size=10, after=4)
doc.add_page_break()

add_heading(doc, "How to Use This Guide", 1)
add_text(doc, "هدف این فایل حفظ‌کردن کلمه‌به‌کلمه نیست. پیام اصلی هر اسلاید، جمله شروع و جمله انتقال را حفظ کنید؛ سپس متن میانی را طبیعی بیان کنید.", rtl_text=True)
for item in [
    "سه بار تمرین کنید: بار اول با متن کامل، بار دوم با Cue Card و بار سوم فقط با پاورپوینت.",
    "سرعت مناسب حدود ۱۱۵ تا ۱۲۵ کلمه در دقیقه است. بعد از هر عدد مهم یک مکث کوتاه داشته باشید.",
    "هنگام توضیح نمودار یا confusion matrix به همان بخش اسلاید اشاره کنید و متن اسلاید را عیناً نخوانید.",
    "در پاسخ به سؤال‌ها بین نتیجه اجراشده، تفسیر و پیشنهاد آینده تفاوت روشن بگذارید.",
]:
    add_text(doc, "• " + item, rtl_text=True, after=4)
add_heading(doc, "Pronunciation Quick Check", 2)
add_text(doc, "MVTec: em-vee-tek  |  Grad-CAM: grad-cam  |  MobileNetV2: mobile-net vee two  |  ROC-AUC: R-O-C A-U-C  |  PR-AUC: P-R A-U-C  |  leakage: lee-kij  |  fine-tuning: fine tuning", size=10.5)
add_callout(doc, "Delivery reminder:", "Do not rush the thresholds or confusion-matrix values. These are the numbers most likely to generate questions.", fill="FFF8E8")
doc.add_page_break()

add_heading(doc, "Part I — Full Slide-by-Slide Script", 1)
add_text(doc, "Suggested total duration: approximately 10–12 minutes. The timing is a guide, not a strict limit.", italic=True, color=MUTED)

for s in slides:
    p = add_heading(doc, f"Slide {s['n']} — {s['title']}", 2)
    add_callout(doc, "زمان و پیام اصلی:", f"{s['time']} — {s['message']}", fill=LIGHT, rtl_text=True)
    add_heading(doc, "Full English script", 3)
    add_text(doc, s["script"], keep=True)
    add_text(doc, "Transition: “" + s["transition"] + "”", italic=True, color=DARK, size=10.5, after=5)
    add_callout(doc, "Cue card:", s["cue"], fill=CALLOUT)

doc.add_page_break()
add_heading(doc, "Part II — Compact Cue Cards", 1)
add_text(doc, "برای تمرین دوم، فقط از این جدول استفاده کنید. ستون عدد کلیدی را دقیق حفظ کنید و بقیه توضیح را با زبان خودتان بگویید.", rtl_text=True)
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.LEFT
set_table_geometry(table, [850, 5660, 2850])
headers = ["Slide", "Keywords", "Number / caution"]
for i,h in enumerate(headers):
    cell=table.rows[0].cells[i]; set_cell_shading(cell,LIGHT); p=cell.paragraphs[0]; r=p.add_run(h); set_font(r,size=9.5,bold=True,color=NAVY)
for s in slides:
    row=table.add_row().cells
    vals=[str(s["n"]), s["cue"], {
        1:"Scope, not metrics",2:"False negatives matter",3:"242 / 54 / 51",4:"Three stages",5:"94.44%; PR-AUC .972",6:"20 layers; 1e-5",7:".152666 / .439",8:"Not causal",9:"94.12%; FN=2",10:"Review missed both FN",11:"Not production-ready"
    }[s["n"]]]
    for i,v in enumerate(vals):
        p=row[i].paragraphs[0]; r=p.add_run(v); set_font(r,size=8.8); row[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
set_table_geometry(table,[850,5660,2850])

doc.add_page_break()
add_heading(doc, "Part III — Likely Questions and Suggested Answers", 1)
add_text(doc, "پاسخ‌ها را کوتاه نگه دارید: ابتدا جواب مستقیم، سپس یک دلیل یا عدد. اگر سؤال خارج از داده‌های پروژه بود، صادقانه آن را به‌عنوان کار آینده معرفی کنید.", rtl_text=True)
for i,(q,a) in enumerate(questions,1):
    add_heading(doc, f"Q{i}. {q}", 3)
    add_text(doc, a, after=7)

doc.add_page_break()
add_heading(doc, "Part IV — Emergency 5-Minute Version", 1)
add_text(doc, "اگر زمان ارائه ناگهان به پنج دقیقه کاهش یافت، برای هر اسلاید تقریباً ۲۰ تا ۳۰ ثانیه صحبت کنید و فقط پیام‌های زیر را بگویید.", rtl_text=True)
for line in short_script:
    num, body = line.split(" — ",1)
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.15
    r=p.add_run(num+" — "); set_font(r,size=10,bold=True,color=DARK)
    r=p.add_run(body); set_font(r,size=10)

doc.add_page_break()
add_heading(doc, "Final Rehearsal Checklist", 2)
for item in [
    "Can I explain why recall matters without saying accuracy is unimportant?",
    "Can I state the split sizes and confirm that the test set was untouched?",
    "Can I explain why the baseline's 100% recall was misleading?",
    "Can I state both locked thresholds and the validation review rate?",
    "Can I read the final confusion matrix correctly: TN 38, FP 1, FN 2, TP 10?",
    "Can I explain honestly why the human-review policy did not catch the two false negatives?",
    "Can I finish with one concrete next step and say the system is not yet production-ready?",
]: add_text(doc, "☐ " + item, after=3, size=10.5)

add_callout(doc, "Last sentence to remember:", "The model is promising, but the missed subtle defects show why explainability, human review, and new holdout validation must remain part of the system design.", fill="FFF8E8")

doc.core_properties.title = "Industrial Surface Defect Detection — Presentation Practice Guide"
doc.core_properties.subject = "MVTec AD Tile final project oral presentation rehearsal"
doc.core_properties.author = "Atousa"
doc.save(OUT)
print(OUT)
